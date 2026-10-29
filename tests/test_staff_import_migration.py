import io, os
from openpyxl import Workbook, load_workbook
from app.extensions import db
from app.models import Product, Category, Tag, Banner, HomepageSection, ProductVariant, Order
from app.services import importer
from .conftest import make_user, login, make_product, make_cat


def xlsx(products, variants=None):
    wb = Workbook(); ws = wb.active; ws.title = 'Products'
    ws.append(importer.PRODUCT_COLS)
    for r in products:
        ws.append(r)
    if variants:
        wv = wb.create_sheet('Variants'); wv.append(importer.VARIANT_COLS)
        for r in variants:
            wv.append(r)
    b = io.BytesIO(); wb.save(b); b.seek(0)
    return b


def test_bulk_import_create_update_dry_run_errors(app):
    rows = [['TSH-1', 'Tee', 20, '', 'published', 'Raven', 'Men > Tops > Tees', '', 'Nice', 'cotton, summer',
             'Material: Cotton; Season: Summer, Winter', 'Color: Black, White; Size: S, M', 7, 'https://x.test/a.jpg'],
            ['CAP-1', 'Cap', 9, 5, 'draft', '', 'Accessories', '', '', 'hat', '', '', 3, ''],
            ['BAD-1', '', 'abc', '', '', '', '', '', '', '', '', '', '', ''],
            ['CAP-1', 'Dup', 1, '', '', '', '', '', '', '', '', '', '', '']]
    variants = [['TSH-1', 'Color=Black; Size=M', 'TSH-1-BM', 2, 25, '', 'yes'],
                ['TSH-1', 'Color=Pink; Size=M', '', 1, '', '', ''], ['NOPE', 'Color=Black', '', 1, '', '', '']]
    dry = importer.import_file(xlsx(rows, variants), 'p.xlsx', apply=False)
    assert Product.query.count() == 0 and Category.query.count() == 0 and dry.created == 2   # nothing saved
    rep = importer.import_file(xlsx(rows, variants), 'p.xlsx', apply=True)
    assert rep.created == 2 and rep.variants_updated == 1
    assert {(s, r) for s, r, _ in rep.errors} == {('Products', 4), ('Products', 5), ('Variants', 3), ('Variants', 4)}
    tee = Product.query.filter_by(sku='TSH-1').one()
    assert len(tee.variants) == 4 and all(v.stock == 7 for v in tee.variants if v.sku != 'TSH-1-BM')
    bm = ProductVariant.query.filter_by(sku='TSH-1-BM').one()
    assert bm.stock == 2 and float(bm.price) == 25 and tee.min_price == 20
    assert [c.name for c in tee.category.ancestors()] == ['Men', 'Tops'] and {t.name for t in tee.tags} == {'cotton', 'summer'}
    assert len(tee.images) == 1 and {a.value for a in tee.attribute_values} == {'Cotton', 'Summer', 'Winter'}
    assert Product.query.filter_by(sku='CAP-1').one().status == 'draft'
    # re-import = update (blank cells keep values, images not duplicated)
    rep2 = importer.import_file(xlsx([['TSH-1', '', 22, '', '', '', '', '', '', '', '', '', '', 'https://x.test/a.jpg']]), 'p.xlsx')
    assert rep2.updated == 1 and not rep2.errors
    tee = Product.query.filter_by(sku='TSH-1').one()
    assert tee.name == 'Tee' and float(tee.price) == 22 and len(tee.images) == 1 and len(tee.variants) == 4
    csv = io.BytesIO(b'sku,name,price,stock\nCSV-1,Csv Item,5,9\n')
    assert importer.import_file(csv, 'x.csv').created == 1
    try:
        importer.import_file(io.BytesIO(b'junk'), 'x.xlsx'); assert False
    except importer.ImportError_:
        pass
    # template and export are valid workbooks that round-trip
    load_workbook(importer.build_template())
    exported = importer.export_products()
    assert importer.import_file(exported, 'p.xlsx', apply=False).errors == []


def test_staff_panel_access_banner_tags_import_and_csrf(app):
    cust = make_user('09120000010'); staff = make_user('09120000002', 'staff'); admin = make_user('09120000001', 'admin')
    anon = app.test_client().get('/staff/products')
    assert anon.status_code == 302 and '/auth/login' in anon.headers['Location']   # anonymous -> login page
    c = app.test_client(); login(c, cust)
    assert c.get('/staff/products').status_code == 403
    assert c.get('/staff/settings').status_code == 403
    c = app.test_client(); login(c, staff)
    for url in ['/staff/', '/staff/products', '/staff/products/new', '/staff/categories', '/staff/categories/new', '/staff/tags',
                '/staff/import', '/staff/import/template.xlsx', '/staff/export/products.xlsx', '/staff/orders', '/staff/settings',
                '/staff/homepage', '/staff/homepage/new', '/staff/banners', '/staff/banners/new', '/staff/coupons',
                '/staff/coupons/new', '/staff/shipping', '/staff/shipping/new']:
        assert c.get(url).status_code == 200, url
    assert c.get('/admin/').status_code == 403 and c.get('/admin/users').status_code == 403
    # categories: nested creation + cycle protection through the UI
    c.post('/staff/categories/new', data={'name': 'Women', 'visible': 'on'})
    w = Category.query.filter_by(name='Women').one()
    c.post('/staff/categories/new', data={'name': 'Dresses', 'parent_id': w.id, 'visible': 'on'})
    d = Category.query.filter_by(name='Dresses').one(); assert d.parent_id == w.id and d.depth == 1
    r = c.post(f'/staff/categories/{w.id}/edit', data={'name': 'Women', 'parent_id': d.id, 'visible': 'on'})
    db.session.refresh(w); assert w.parent_id is None and b'descendants' in r.data
    c.post('/staff/categories/new', data={'name': 'Picks', 'is_internal': 'on'})
    assert Category.query.filter_by(name='Picks').one().public_effective is False
    # product with tags (existing + new), staying hidden from public until published
    c.post('/staff/tags', data={'names': 'summer, cotton'})
    tag_id = Tag.query.filter_by(slug='summer').one().id
    r = c.post('/staff/products/new', data={'name': 'Tee', 'price': '10', 'stock': '4', 'status': 'draft', 'sku': 'T-1',
                                           'category_id': d.id, 'tag_ids': tag_id, 'new_tags': 'sale, Summer'})
    p = Product.query.filter_by(sku='T-1').one()
    assert {t.name for t in p.tags} == {'summer', 'sale'} and p.stock == 4
    assert app.test_client().get(f'/product/{p.slug}').status_code == 404
    c.post(f'/staff/products/{p.id}/status', data={'status': 'published'})
    assert app.test_client().get(f'/product/{p.slug}').status_code == 200
    bad = c.post('/staff/products/new', data={'name': 'X', 'price': '10', 'sale_price': '20', 'stock': '-1', 'sku': 'T-1'})
    assert bad.status_code == 200 and Product.query.count() == 1                       # validation blocks bad saves
    # banner + section through the panel
    c.post('/staff/banners/new', data={'image_url': 'https://x.test/b.jpg', 'title': 'Big sale', 'link_type': 'product',
                                      'product_ref': 'T-1', 'is_active': 'on', 'position': '1'})
    b = Banner.query.one(); assert b.product_id == p.id
    assert 'Big sale' in app.test_client().get('/').get_data(as_text=True)
    c.post('/staff/homepage/new', data={'source_category_id': d.id, 'product_limit': '10', 'auto_scroll': 'on',
                                        'scroll_interval': '4', 'is_active': 'on', 'visible': 'on', 'include_subcategories': 'on'})
    s = HomepageSection.query.filter(HomepageSection.source_category_id == d.id).one()
    assert s.auto_scroll and s.scroll_interval == 4 and s.display_title == 'Dresses'
    home = app.test_client().get('/').get_data(as_text=True)
    assert 'data-autoscroll="1"' in home and 'data-interval="4000"' in home
    # import through the UI
    r = c.post('/staff/import', data={'mode': 'apply', 'file': (xlsx([['UI-1', 'Ui Item', 3, '', 'published', '', '', '', '', '', '', '', 2, '']]), 'a.xlsx')},
               content_type='multipart/form-data')
    assert r.status_code == 200 and Product.query.filter_by(sku='UI-1').count() == 1
    # product delete: with order history it is archived, not removed
    from .test_orders_payments import add_cart, place
    u = make_user('09120000020'); add_cart(u, p); place(u)
    c.post(f'/staff/products/{p.id}/delete')
    assert db.session.get(Product, p.id).status == 'archived'
    # admin-only pages
    a = app.test_client(); login(a, admin)
    assert a.get('/admin/users').status_code == 200 and a.get('/admin-panel', follow_redirects=True).status_code == 200 and a.get('/admin-panel/admin_tags/').status_code == 200 \
        and a.get('/admin-panel/admin_products/').status_code == 200 and a.get('/admin-panel/admin_products/new/').status_code in (302, 403, 404)


def test_csrf_enforced_when_enabled(app):
    app.config['WTF_CSRF_ENABLED'] = True
    c = app.test_client()
    assert c.post('/auth/login', data={'phone': '0912', 'password': 'x'}).status_code == 400


def test_production_config_rejects_unsafe_settings():
    import config
    cfg = type('P', (config.ProductionConfig,), {'SECRET_KEY': 'dev-secret-change-me', 'SMS_PROVIDER': 'mock', 'SITE_URL': 'http://x'})
    try:
        cfg.validate(); assert False
    except RuntimeError as e:
        assert 'SECRET_KEY' in str(e) and 'mock' in str(e)


def test_upload_validation(app, tmp_path):
    from app.uploads import save_image, UploadError, validate_image_url
    from werkzeug.datastructures import FileStorage
    import pytest
    with app.test_request_context():
        with pytest.raises(UploadError):
            save_image(FileStorage(io.BytesIO(b'<svg onload=alert(1)>'), filename='x.png'))
        from PIL import Image
        b = io.BytesIO(); Image.new('RGB', (4, 4)).save(b, 'PNG'); b.seek(0)
        url = save_image(FileStorage(b, filename='evil.php'))
        assert url.startswith('/media/products/') and url.endswith('.png')
        with pytest.raises(UploadError):
            validate_image_url('javascript:alert(1)')


def test_migrations_preserve_legacy_data(tmp_path):
    """Build the ORIGINAL schema at revision 0001 with rows, upgrade to head, verify data was mapped."""
    import sqlite3
    from flask_migrate import upgrade, stamp
    from app import create_app
    import config
    db_path = tmp_path / 'legacy.db'
    cfg = type('C', (config.TestingConfig,), {'SQLALCHEMY_DATABASE_URI': f'sqlite:///{db_path}', 'SQLALCHEMY_ENGINE_OPTIONS': {}})
    app = create_app(cfg)
    with app.app_context():
        upgrade(revision='0001_baseline')
        c = sqlite3.connect(db_path)
        c.executescript("""
        insert into users(phone,name,password_hash,role,created_at) values('09120000003','Cust','x','customer','2025-01-01');
        insert into categories(name,slug,visible,image_url) values('Clothing','clothing',1,''),('Hidden','hidden',0,'');
        insert into products(name,slug,description,price,stock,image_url,category_id,created_at,is_active) values
          ('Shirt','shirt','d',20.5,7,'http://i/x.jpg',1,'2025-01-02',1),('Old','old','d',5,0,'',2,'2025-01-03',0);
        insert into cart_items(user_id,product_id,quantity) values(1,1,2);
        insert into orders(user_id,status,total_amount,shipping_address,created_at) values(1,'processing',41,'addr','2025-02-01'),(1,'delivered',5,'a2','2025-02-02');
        insert into order_items(order_id,product_id,quantity,price_at_purchase) values(1,1,2,20.5),(2,2,1,5);""")
        c.commit(); c.close()
        upgrade()
        c = sqlite3.connect(db_path)
        q = lambda s: c.execute(s).fetchall()
        assert q("select name,status,min_price from products order by id") == [('Shirt', 'published', 20.5), ('Old', 'draft', 5)]
        assert q("select v.stock, v.option_key from product_variants v join products p on p.id=v.product_id where p.name='Shirt'") == [(7, '')]
        assert q("select url,is_primary from product_images") == [('http://i/x.jpg', 1)]
        assert q("select path,public_effective from categories order by id") == [('/1/', 1), ('/2/', 0)]
        assert q("select variant_id,quantity from cart_items") == [(1, 2)]
        assert q("select status,subtotal from orders order by id") == [('preparing', 41), ('received', 5)]
        assert q("select product_name,stock_state,variant_id from order_items order by id") == [('Shirt', 'deducted', 1), ('Old', 'deducted', 2)]
        assert q("select count(*) from users where phone_verified=1") == [(1,)]
        assert q("select title from homepage_sections") == [('New arrivals',)]
        c.close()
        from sqlalchemy import text
        from app.extensions import db
        db.session.remove(); db.engine.dispose()


def test_product_form_renders_with_data(app):
    """Regression: GET new/edit crashed with UndefinedError once categories/tags/products existed."""
    from .test_orders_payments import add_cart
    staff = make_user('09120000002', 'staff'); c = app.test_client(); login(c, staff)
    cat = make_cat('Men'); sub = make_cat('Tops', cat); internal = make_cat('Picks', is_internal=True)
    p = make_product('Tee', 10, [sub, internal], options='Color: Red, Blue\nSize: M', tags='a, b', attributes='Material: Cotton',
                     new_image_urls='https://x.test/a.jpg')
    simple = make_product('Cap', 5, [cat], tags='a')
    for url in ('/staff/products/new', f'/staff/products/{p.id}/edit', f'/staff/products/{simple.id}/edit', '/staff/categories',
                f'/staff/categories/{cat.id}/edit', '/staff/categories/new?parent=%d' % cat.id):
        r = c.get(url); assert r.status_code == 200, url
    html = c.get(f'/staff/products/{p.id}/edit').get_data(as_text=True)
    assert 'Color: Red, Blue' in html and 'name="tag_ids"' in html and 'Material: Cotton' in html
    r = c.post(f'/staff/products/{p.id}/edit', data={'name': 'Tee', 'price': '10', 'status': 'published',
               'options': 'Color: Red, Blue\nSize: M', 'category_id': sub.id, 'tag_ids': [], 'new_tags': 'c'})
    assert r.status_code == 302
