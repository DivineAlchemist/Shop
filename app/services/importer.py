"""Bulk product import/export (Excel .xlsx or CSV).

Products sheet (one row per product, upsert by SKU):
  sku*, name*, price*, sale_price, status, brand, category ("Women > Tops > Shirts"),
  extra_categories (";" separated paths), description, tags ("a, b"), attributes ("Material: Cotton; Season: Summer, Winter"),
  options ("Color: Blue, Black; Size: M, L"), stock, image_urls (comma/space separated URLs)
Variants sheet (optional): product_sku*, options* ("Color=Blue; Size=M"), sku, stock, price, sale_price, active
Blank cells on an existing product mean "leave unchanged".
"""
import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal
from werkzeug.datastructures import MultiDict

from ..extensions import db
from ..models import Product, ProductVariant, Category
from ..utils import to_decimal, unique_slug
from .catalog import rebuild_subtree
from .products import save_product, options_to_text, combo_key

MAX_ROWS = 5000
PRODUCT_COLS = ['sku', 'name', 'price', 'sale_price', 'status', 'brand', 'category', 'extra_categories',
                'description', 'tags', 'attributes', 'options', 'stock', 'image_urls']
VARIANT_COLS = ['product_sku', 'options', 'sku', 'stock', 'price', 'sale_price', 'active']
EXAMPLES = [
    ['TSH-001', 'Classic T-Shirt', 19.9, '', 'published', 'Raven', 'Men > Tops > T-Shirts', 'Summer', 'Soft cotton tee.',
     'cotton, summer', 'Material: Cotton; Fit: Regular', 'Color: Black, White; Size: S, M, L', 10, 'https://example.com/tee.jpg'],
    ['CAP-001', 'Baseball Cap', 12, 9.5, 'draft', 'Raven', 'Accessories > Hats', '', '', 'accessory', '', '', 25, ''],
]
VARIANT_EXAMPLES = [['TSH-001', 'Color=Black; Size=M', 'TSH-001-BLK-M', 4, '', '', 'yes']]


class ImportError_(Exception):
    pass


@dataclass
class Report:
    created: int = 0
    updated: int = 0
    variants_updated: int = 0
    errors: list = field(default_factory=list)   # (sheet, row_number, message)
    dry_run: bool = False

    @property
    def ok(self):
        return not self.errors


# ------------------------------------------------------------------ reading
def _norm(h):
    return str(h or '').strip().lower().replace(' ', '_')


def _cell(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_rows(stream, filename):
    """Return {'products': [dict], 'variants': [dict]} with 1-based sheet row numbers in '_row'."""
    name = (filename or '').lower()
    out = {'products': [], 'variants': []}
    if name.endswith('.csv'):
        text = stream.read().decode('utf-8-sig', errors='replace')
        reader = csv.reader(io.StringIO(text))
        sheets = {'products': list(reader)}
    elif name.endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(stream, read_only=True, data_only=True)
        except Exception:
            raise ImportError_('The file is not a valid .xlsx workbook.')
        sheets = {}
        for ws in wb.worksheets:
            key = ws.title.strip().lower()
            if key in ('products', 'variants'):
                sheets[key] = [list(r) for r in ws.iter_rows(values_only=True)]
        if 'products' not in sheets:
            first = wb.worksheets[0]
            sheets['products'] = [list(r) for r in first.iter_rows(values_only=True)]
    else:
        raise ImportError_('Upload a .xlsx or .csv file.')
    for key, rows in sheets.items():
        if not rows:
            continue
        headers = [_norm(h) for h in rows[0]]
        if len(rows) - 1 > MAX_ROWS:
            raise ImportError_(f'Too many rows (max {MAX_ROWS}).')
        for i, r in enumerate(rows[1:], start=2):
            d = {h: _cell(r[j]) if j < len(r) else '' for j, h in enumerate(headers) if h}
            if any(d.values()):
                d['_row'] = i
                out[key].append(d)
    return out


# ------------------------------------------------------------------ helpers
def ensure_category_path(path):
    parent = None
    for seg in [s.strip() for s in path.split('>') if s.strip()]:
        q = Category.query.filter(db.func.lower(Category.name) == seg.lower(),
                                  Category.parent_id == (parent.id if parent else None))
        cat = q.first()
        if cat is None:
            cat = Category(name=seg[:80], slug=unique_slug(Category, seg), parent=parent, visible=True)
            db.session.add(cat)
            db.session.flush()
            rebuild_subtree(cat)
        parent = cat
    return parent


def _split_images(text):
    return [u for u in text.replace(',', ' ').replace(';', ' ').split() if u]


def _build_form(row, existing):
    form = MultiDict()

    def pick(key, current=''):
        v = row.get(key, '')
        return v if v != '' else (current if existing else '')
    form['name'] = pick('name', existing.name if existing else '')
    form['price'] = pick('price', str(existing.price) if existing else '')
    form['sale_price'] = pick('sale_price', str(existing.sale_price) if existing and existing.sale_price is not None else '')
    form['status'] = pick('status', existing.status if existing else 'draft') or 'draft'
    form['brand'] = pick('brand', existing.brand if existing else '')
    form['description'] = pick('description', existing.description if existing else '')
    form['sku'] = row.get('sku', '')
    form['stock'] = row.get('stock') or (str(sum(v.stock for v in existing.variants)) if existing and not existing.options else '0')
    opts = row.get('options', '').replace(';', '\n')
    form['options'] = opts if opts else (options_to_text(existing.options) if existing else '')
    form['tags'] = row.get('tags', '').replace(';', ',') if row.get('tags') else (
        ', '.join(t.name for t in existing.tags) if existing else '')
    attrs = row.get('attributes', '').replace(';', '\n')
    form['attributes'] = attrs if attrs else (
        '\n'.join(f'{k}: {", ".join(v)}' for k, v in _attrs_of(existing).items()) if existing else '')
    cats = []
    if row.get('category'):
        main = ensure_category_path(row['category'])
        if main:
            form['category_id'] = str(main.id)
            cats.append(main.id)
    elif existing:
        if existing.category_id:
            form['category_id'] = str(existing.category_id)
        cats += [c.id for c in existing.categories]
    for p in [x for x in row.get('extra_categories', '').split(';') if x.strip()]:
        c = ensure_category_path(p)
        if c:
            cats.append(c.id)
    for cid in dict.fromkeys(cats):
        form.add('category_ids', str(cid))
    have = {i.url for i in existing.images} if existing else set()
    for u in _split_images(row.get('image_urls', '')):
        if u not in have:
            form.add('new_image_urls', u)
    if existing and 'new_image_urls' in form:
        form['new_image_urls'] = '\n'.join(form.getlist('new_image_urls'))
    elif 'new_image_urls' in form:
        form['new_image_urls'] = '\n'.join(form.getlist('new_image_urls'))
    return form


def _attrs_of(product):
    out = {}
    for pav in product.attribute_values:
        out.setdefault(pav.attribute.name, []).append(pav.value)
    return out


def _parse_variant_options(text):
    pairs = []
    for part in text.split(';'):
        if '=' in part:
            g, v = part.split('=', 1)
            pairs.append((g.strip(), v.strip()))
    return pairs


# ------------------------------------------------------------------ import
def import_file(stream, filename, apply=True):
    report = Report(dry_run=not apply)
    data = _read_rows(stream, filename)
    seen = set()
    for row in data['products']:
        n = row['_row']
        sku = row.get('sku', '')
        if not sku:
            report.errors.append(('Products', n, 'SKU is required.'))
            continue
        if sku in seen:
            report.errors.append(('Products', n, f'Duplicate SKU "{sku}" in file.'))
            continue
        seen.add(sku)
        try:
            with db.session.begin_nested():
                existing = Product.query.filter_by(sku=sku).first()
                form = _build_form(row, existing)
                product, errs = save_product(existing, form, MultiDict())
                if errs:
                    raise ValueError('; '.join(errs))
                if existing is None and row.get('stock', '') != '' and product.options:
                    for v in product.variants:
                        v.stock = int(row['stock']) if row['stock'].isdigit() else 0
                    product.refresh_price_cache()
                report.created += existing is None
                report.updated += existing is not None
        except Exception as e:  # row-level failure never aborts the whole file
            report.errors.append(('Products', n, str(e)[:300]))
    for row in data['variants']:
        n = row['_row']
        try:
            with db.session.begin_nested():
                product = Product.query.filter_by(sku=row.get('product_sku', '')).first()
                if product is None:
                    raise ValueError(f'Unknown product SKU "{row.get("product_sku", "")}".')
                combo = []
                for g, label in _parse_variant_options(row.get('options', '')):
                    vals = (product.options or {}).get(g)
                    vid = next((k for k, l in (vals or {}).items() if l.lower() == label.lower()), None)
                    if vid is None:
                        raise ValueError(f'Option "{g}={label}" does not exist on {product.sku}.')
                    combo.append((g, vid))
                variant = ProductVariant.query.filter_by(product_id=product.id, option_key=combo_key(combo)).first()
                if variant is None:
                    raise ValueError('No variant matches those options.')
                if row.get('sku'):
                    clash = ProductVariant.query.filter(ProductVariant.sku == row['sku'], ProductVariant.id != variant.id).first()
                    if clash:
                        raise ValueError(f'Variant SKU "{row["sku"]}" is already in use.')
                    variant.sku = row['sku'][:80]
                if row.get('stock', '') != '':
                    if not row['stock'].isdigit():
                        raise ValueError('Stock must be a whole number.')
                    variant.stock = max(int(row['stock']), variant.reserved or 0)
                for col in ('price', 'sale_price'):
                    if row.get(col, '') != '':
                        d = to_decimal(row[col])
                        if d is None or d < 0:
                            raise ValueError(f'Invalid {col}.')
                        setattr(variant, col, d)
                if row.get('active', '') != '':
                    variant.is_active = row['active'].lower() in ('1', 'yes', 'true', 'y')
                db.session.flush()
                product.refresh_price_cache()
                report.variants_updated += 1
        except Exception as e:
            report.errors.append(('Variants', n, str(e)[:300]))
    if apply:
        db.session.commit()
    else:
        db.session.rollback()
    return report


# ------------------------------------------------------------------ templates / export
def _workbook(product_rows, variant_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = 'Products'
    ws.append(PRODUCT_COLS)
    for r in product_rows:
        ws.append(r)
    wv = wb.create_sheet('Variants')
    wv.append(VARIANT_COLS)
    for r in variant_rows:
        wv.append(r)
    wh = wb.create_sheet('Help')
    for line in __doc__.strip().splitlines():
        wh.append([line])
    for sheet in (ws, wv):
        for c in sheet[1]:
            c.font = Font(bold=True)
        sheet.freeze_panes = 'A2'
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def build_template():
    return _workbook(EXAMPLES, VARIANT_EXAMPLES)


def export_products(limit=MAX_ROWS):
    prows, vrows = [], []
    for p in Product.query.order_by(Product.id).limit(limit).all():
        cat = ' > '.join([c.name for c in p.category.ancestors()] + [p.category.name]) if p.category else ''
        prows.append([p.sku or '', p.name, float(p.price), float(p.sale_price) if p.sale_price is not None else '',
                      p.status, p.brand, cat, '', p.description, ', '.join(t.name for t in p.tags),
                      '; '.join(f'{k}: {", ".join(v)}' for k, v in _attrs_of(p).items()),
                      options_to_text(p.options).replace('\n', '; '),
                      sum(v.stock for v in p.variants) if not p.options else '',
                      ' '.join(i.url for i in p.images)])
        if p.options:
            for v in p.variants:
                vrows.append([p.sku or '', '; '.join(f'{g}={l}' for g, l in v.options_labels.items()), v.sku or '',
                              v.stock, float(v.price) if v.price is not None else '',
                              float(v.sale_price) if v.sale_price is not None else '', 'yes' if v.is_active else 'no'])
    return _workbook(prows, vrows)
